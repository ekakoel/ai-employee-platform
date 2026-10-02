from io import BytesIO
from pathlib import Path
import re

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

import fitz
import pytesseract
from PIL import Image


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".txt",
}


MIN_MEANINGFUL_CHARACTERS = 30


def normalize_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def normalize_text(text: str) -> str:
    if not text:
        return ""

    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    lines = [
        line.rstrip()
        for line in text.split("\n")
    ]

    normalized_lines = []
    blank_count = 0

    for line in lines:
        if not line.strip():
            blank_count += 1

            if blank_count <= 2:
                normalized_lines.append("")

            continue

        blank_count = 0
        normalized_lines.append(line)

    text = "\n".join(normalized_lines)

    text = re.sub(
        r"[ \t]{3,}",
        "  ",
        text,
    )

    return text.strip()


def has_meaningful_text(text: str) -> bool:
    meaningful_characters = sum(
        character.isalnum()
        for character in text
    )

    return meaningful_characters >= MIN_MEANINGFUL_CHARACTERS


def extract_pdf_with_pypdf(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))

    if not reader.pages:
        raise ValueError("PDF contains no pages")

    extracted_pages: list[str] = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        candidates: list[str] = []

        try:
            layout_text = page.extract_text(
                extraction_mode="layout"
            )

            if layout_text:
                candidates.append(layout_text)

        except Exception:
            pass

        try:
            standard_text = page.extract_text()

            if standard_text:
                candidates.append(standard_text)

        except Exception:
            pass

        if not candidates:
            continue

        best_text = max(
            candidates,
            key=lambda value: (
                len(value.strip()),
                sum(
                    character.isalnum()
                    for character in value
                ),
            ),
        )

        best_text = normalize_text(best_text)

        if best_text:
            extracted_pages.append(
                f"[Page {page_number}]\n{best_text}"
            )

    return normalize_text(
        "\n\n".join(extracted_pages)
    )


def extract_pdf_with_ocr(data: bytes) -> str:
    """
    Render PDF pages as images and run OCR.

    This is the fallback for scanned/image-based PDFs.
    """

    pdf = fitz.open(
        stream=data,
        filetype="pdf",
    )

    if pdf.page_count == 0:
        pdf.close()
        raise ValueError("PDF contains no pages")

    extracted_pages: list[str] = []

    try:
        for page_number in range(pdf.page_count):
            page = pdf.load_page(page_number)

            # 200 DPI is a reasonable balance between
            # OCR quality and processing cost.
            matrix = fitz.Matrix(
                200 / 72,
                200 / 72,
            )

            pixmap = page.get_pixmap(
                matrix=matrix,
                alpha=False,
            )

            image_bytes = pixmap.tobytes(
                "png"
            )

            image = Image.open(
                BytesIO(image_bytes)
            )

            text = pytesseract.image_to_string(
                image,
                config="--psm 6",
            )

            text = normalize_text(text)

            if text:
                extracted_pages.append(
                    f"[Page {page_number + 1}]\n{text}"
                )

    finally:
        pdf.close()

    return normalize_text(
        "\n\n".join(extracted_pages)
    )


def extract_pdf(data: bytes) -> str:
    """
    PDF extraction strategy:

    1. Try normal PDF text extraction.
    2. If insufficient, fallback to OCR.
    3. Fail only if both approaches produce
       insufficient meaningful text.
    """

    text = extract_pdf_with_pypdf(data)

    if has_meaningful_text(text):
        return text

    ocr_text = extract_pdf_with_ocr(data)

    if has_meaningful_text(ocr_text):
        return ocr_text

    raise ValueError(
        "PDF text extraction and OCR both produced "
        "insufficient text. The PDF may be corrupted, "
        "protected, or contain content that requires "
        "specialized table/image extraction."
    )


def extract_docx(data: bytes) -> str:
    document = Document(BytesIO(data))

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    table_lines: list[str] = []

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            cells = [
                cell
                for cell in cells
                if cell
            ]

            if cells:
                table_lines.append(
                    " | ".join(cells)
                )

    result_parts = []

    if paragraphs:
        result_parts.append(
            "\n".join(paragraphs)
        )

    if table_lines:
        result_parts.append(
            "\n".join(table_lines)
        )

    return normalize_text(
        "\n\n".join(result_parts)
    )


def extract_xlsx(data: bytes) -> str:
    workbook = load_workbook(
        filename=BytesIO(data),
        read_only=True,
        data_only=True,
    )

    sections: list[str] = []

    for worksheet in workbook.worksheets:
        sections.append(
            f"[Sheet: {worksheet.title}]"
        )

        for row in worksheet.iter_rows(
            values_only=True
        ):
            values = []

            for value in row:
                if value is None:
                    continue

                text = str(value).strip()

                if text:
                    values.append(text)

            if values:
                sections.append(
                    " | ".join(values)
                )

    workbook.close()

    return normalize_text(
        "\n".join(sections)
    )


def extract_txt(data: bytes) -> str:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode(
            "cp1252",
            errors="replace",
        )

    return normalize_text(text)


def extract_text(
    filename: str,
    data: bytes,
) -> str:
    extension = normalize_extension(
        filename
    )

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported document type: {extension}"
        )

    if extension == ".pdf":
        return extract_pdf(data)

    if extension == ".docx":
        return extract_docx(data)

    if extension == ".xlsx":
        return extract_xlsx(data)

    if extension == ".txt":
        return extract_txt(data)

    raise ValueError(
        f"Unsupported document type: {extension}"
    )